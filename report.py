"""Build <data>/out/summary.xlsx and summary.csv from the saved khata JSON files."""
import csv
import json
from collections import defaultdict
from pathlib import Path

import config
from names import TARGETS, match_ror, should_download

OUT = config.OUT
JSON_DIR = OUT / "json"
PDF_DIR = OUT / "pdf"


def gatas(ror):
    return "; ".join(f"{k.get('khasra_no')} ({k.get('area')} हे०)" for k in ror.get("khasra") or [])


def mutations(ror):
    rows = ((ror.get("cnameResponse1") or {}).get("data")) or []
    return "; ".join(f"{m.get('transfer_type_desc')} / {m.get('order_type_desc')} / {m.get('order_date')}" for m in rows)


def main():
    matches, all_khatas = [], []
    per_person = defaultdict(list)
    for f in sorted(JSON_DIR.glob("khata_*.json")):
        saved = json.loads(f.read_text(encoding="utf-8"))
        khata, ror = saved["khata_number"], saved["ror"]
        hits = match_ror(ror)
        pdf = PDF_DIR / f"khata_{khata}.pdf"
        land = "; ".join(d.get("land_type") or "" for d in ror.get("landDesc") or [])
        all_khatas.append({"khata": khata, "gatas": gatas(ror), "land_type": land,
                           "target_hits": len(hits), "pdf": pdf.name if pdf.exists() else ""})
        for h in hits:
            row = {
                "khata": khata, "gatas (khasra, area)": gatas(ror), "land_type": land,
                "matched_target": f"{h['target']} / {h['target_father']}",
                "name_on_record": h["name"], "father_on_record": h["father"], "address": h["address"],
                "hissa": h["hissa"], "role": h["source"], "match": h["strength"],
                "downloaded": "yes" if should_download(hits) and pdf.exists() else "no",
                "pdf": str(pdf.relative_to(OUT)) if pdf.exists() else "",
                "mutations": mutations(ror),
            }
            matches.append(row)
            per_person[row["matched_target"]].append(f"{khata} [{h['source']},{h['strength']}]")

    order = {"exact": 0, "fuzzy": 1, "name_only": 2}
    matches.sort(key=lambda r: (order[r["match"]], r["khata"]))
    if matches:
        with (OUT / "summary.csv").open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(matches[0]))
            w.writeheader()
            w.writerows(matches)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("openpyxl missing; wrote CSV only")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    if matches:
        ws.append(list(matches[0]))
        for r in matches:
            ws.append(list(r.values()))
    ws2 = wb.create_sheet("Per person")
    ws2.append(["target (name / father)", "khatas [role, match]"])
    for t_name, t_father in TARGETS:
        key = f"{t_name} / {t_father}"
        ws2.append([key, ", ".join(per_person.get(key, [])) or "—"])
    ws3 = wb.create_sheet("All khatas")
    ws3.append(list(all_khatas[0]) if all_khatas else ["khata"])
    for r in all_khatas:
        ws3.append(list(r.values()))
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    wb.save(OUT / "summary.xlsx")
    print(f"{len(all_khatas)} khatas scanned, {len(matches)} target hits, "
          f"{sum(1 for r in all_khatas if r['pdf'])} PDFs → {OUT / 'summary.xlsx'}, {OUT / 'summary.csv'}")


if __name__ == "__main__":
    main()
